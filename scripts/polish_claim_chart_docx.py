#!/usr/bin/env python
"""Check claim-chart formatting; perform only explicitly requested, narrow repairs.

Assumes a summary table followed by a main claim table, with the company in
column two and image-first evidence. Default operation is read-only. Repairs write a separate DOCX,
preserve non-target runs/package parts, and never invent correspondence text.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run


HEREIN = re.compile(r'Herein,\s*[“"](?P<claim>.+?)[”"]\s+corresponds to\s+(?P<mapping>.+)', re.S)
REF = re.compile(r'^\s*\[Ref-(\d+)\]')
INLINE_REF = re.compile(r'\[Ref-\d+[^\]]*\]')
URL = re.compile(r'https?://[^\s<>]+')


@dataclass(frozen=True)
class Issue:
    code: str
    message: str

    def __str__(self):
        return f'[{self.code}] {self.message}'


class UnsafeRepair(ValueError):
    """The helper cannot preserve meaning/structure while making this repair."""


def supported_layout(document):
    return len(document.tables) >= 2 and all(
        len(table.rows) and all(len(row.cells) >= 2 for row in table.rows)
        for table in document.tables[:2]
    )


def run_spans(paragraph):
    """Include hyperlink runs, but exclude paragraphs inside drawing/text boxes."""
    offset = 0
    for element in paragraph._p.xpath('.//w:r'):
        parent = element.getparent()
        while parent is not None and parent.tag != qn('w:p'):
            parent = parent.getparent()
        if parent is not paragraph._p:
            continue
        run = Run(element, paragraph)
        length = len(run.text)
        if length:
            yield offset, offset + length, run
            offset += length


def paragraph_text(paragraph):
    return ''.join(run.text for _, _, run in run_spans(paragraph))


def effective(run, paragraph, prop):
    direct = getattr(run.font, prop)
    if direct is not None:
        return direct
    for style in (run.style, paragraph.style):
        seen = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            value = getattr(style.font, prop)
            if value is not None:
                return value
            style = style.base_style
    tag = 'b' if prop == 'bold' else 'i'
    defaults = paragraph.part.document.styles.element.xpath(
        f'./w:docDefaults/w:rPrDefault/w:rPr/w:{tag}'
    )
    return bool(defaults and defaults[0].get(qn('w:val'), '1') not in ('0', 'false', 'off'))


def span_has(paragraph, start, end, prop):
    found = False
    for left, right, run in run_spans(paragraph):
        if left >= end or right <= start:
            continue
        sample = run.text[max(start - left, 0):min(end - left, right - left)]
        if sample.strip():
            found = True
            if not effective(run, paragraph, prop):
                return False
    return found


def format_span(paragraph, start, end, **properties):
    """Split only plain text runs at the requested boundaries; keep rPr and pPr."""
    changed = False
    for left, right, run in list(run_spans(paragraph)):
        begin, finish = max(start, left) - left, min(end, right) - left
        if begin >= finish or all(effective(run, paragraph, key) == value for key, value in properties.items()):
            continue
        if begin == 0 and finish == len(run.text):
            targets = [run]
        else:
            if any(child.tag not in (qn('w:rPr'), qn('w:t')) for child in run._r):
                raise UnsafeRepair('A marker crosses a run containing non-text content; format it manually.')
            text = run.text
            targets = []
            for a, b in zip([0, begin, finish], [begin, finish, len(text)]):
                if a == b:
                    continue
                element = deepcopy(run._r)
                clone = Run(element, paragraph)
                clone.text = text[a:b]
                run._r.addprevious(element)
                if a == begin and b == finish:
                    targets.append(clone)
            run._r.getparent().remove(run._r)
        for target in targets:
            for key, value in properties.items():
                setattr(target.font, key, value)
        changed = True
    return int(changed)


def comment_paragraphs(cell):
    paragraphs = [p for p in cell.paragraphs if paragraph_text(p).strip()]
    for i, paragraph in enumerate(paragraphs):
        text = paragraph_text(paragraph).strip()
        if REF.match(text):
            return paragraphs[:i]
        if text.startswith('Herein,') and text.endswith(']'):
            return paragraphs[:i + 1]
    return paragraphs


def format_references_label(document):
    changed = 0
    for paragraph in document.tables[0].cell(0, 1).paragraphs:
        text = paragraph_text(paragraph)
        if text.strip() == 'References:':
            start = text.index('References:')
            changed += format_span(paragraph, start, start + len('References:'), bold=True)
    return changed


def format_comment_opening(cell):
    paragraphs = comment_paragraphs(cell)
    if not paragraphs:
        return 0
    paragraph = paragraphs[0]
    text = paragraph_text(paragraph)
    if not text.lstrip().startswith('[Comment:'):
        return 0
    start = text.index('[Comment:')
    return format_span(paragraph, start, start + len('[Comment:'), bold=True, italic=True)


def format_comment_closing(cell):
    paragraphs = comment_paragraphs(cell)
    if not paragraphs:
        return 0
    paragraph = paragraphs[-1]
    text = paragraph_text(paragraph).rstrip()
    if not text.lstrip().startswith('Herein,') or not text.endswith(']'):
        return 0
    return format_span(paragraph, len(text) - 1, len(text), bold=True, italic=True)


def evidence_events(cell):
    events = []
    for paragraph in cell.paragraphs:
        image = bool(paragraph._p.xpath('.//w:drawing | .//w:pict'))
        reference = bool(REF.match(paragraph_text(paragraph)))
        if image or reference:
            events.append(('MIXED' if image and reference else 'IMG' if image else 'REF', paragraph))
    return events


def image_first(events):
    kinds = [kind for kind, _ in events]
    if 'MIXED' in kinds:
        return False
    if 'IMG' not in kinds:
        return True  # Supporting references alone do not fabricate visual evidence.
    return kinds[0] == 'IMG' and all(
        i + 1 < len(kinds) and kinds[i + 1] == 'REF'
        for i, kind in enumerate(kinds) if kind == 'IMG'
    )


def reorder_evidence(cell):
    """Explicit URL-first repair. Validate every pair before moving any node."""
    events = evidence_events(cell)
    if image_first(events):
        return 0
    kinds = [kind for kind, _ in events]
    if not kinds or len(kinds) % 2 or kinds != ['REF', 'IMG'] * (len(kinds) // 2):
        raise UnsafeRepair('Evidence is mixed or ambiguous; pair each screenshot with its source manually.')
    for i in range(0, len(events), 2):
        ref, image = events[i][1]._p, events[i + 1][1]._p
        sibling = ref.getnext()
        while sibling is not image:
            if sibling is None or sibling.tag != qn('w:p') or sibling.xpath('.//w:t | .//w:drawing | .//w:pict | .//w:br | .//w:fldChar'):
                raise UnsafeRepair('A caption or other content separates a URL/image pair; reorder it manually.')
            sibling = sibling.getnext()
    for i in range(0, len(events), 2):
        events[i][1]._p.addprevious(events[i + 1][1]._p)
    return len(events) // 2


def reference_lines(cell):
    lines = []
    for paragraph in cell.paragraphs:
        text = paragraph_text(paragraph)
        match = REF.match(text)
        if not match:
            continue
        urls = []
        for link in paragraph._p.xpath('.//w:hyperlink'):
            relationship = paragraph.part.rels.get(link.get(qn('r:id')))
            if relationship is not None and relationship.is_external:
                urls.append(relationship.target_ref)
        urls.extend(URL.findall(text))
        urls = list(dict.fromkeys(url.split('#', 1)[0] for url in urls if url.startswith(('https://', 'http://'))))
        lines.append((int(match.group(1)), urls))
    return lines


def summary_prose(document):
    texts = [paragraph_text(p).strip() for p in document.tables[0].cell(0, 1).paragraphs if paragraph_text(p).strip()]
    return [text for text in texts[1:] if text != 'References:' and not REF.match(text)]


def collect_issues(document, template=None):
    issues = []

    def report(code, message):
        issues.append(Issue(code, message))

    if not supported_layout(document):
        return [Issue('LAYOUT', 'Expected a summary and main table with at least two columns in every row.')]
    if template is not None and not supported_layout(template):
        return [Issue('TEMPLATE_LAYOUT', 'The reference does not use the supported two-table layout.')]
    template_comments = [
        p for row in template.tables[1].rows for p in comment_paragraphs(row.cells[1])
    ] if template is not None else []
    allow_inline = any(INLINE_REF.search(paragraph_text(p)) for p in template_comments)
    template_hereins = [(p, HEREIN.search(paragraph_text(p))) for p in template_comments]
    require_emphasis = template is None or any(
        match and span_has(p, *match.span('claim'), 'bold') for p, match in template_hereins
    )
    if summary_prose(document) and not (template is not None and summary_prose(template)):
        report('SUMMARY_PROSE', 'Unexpected opening/summary prose; the reference proceeds from the company heading to References.')

    summary = document.tables[0].cell(0, 1)
    labels = [p for p in summary.paragraphs if paragraph_text(p).strip() == 'References:']
    if len(labels) != 1:
        report('REFERENCES_LABEL', 'Expected exactly one References: label in the summary.')
    elif not span_has(labels[0], 0, len(paragraph_text(labels[0])), 'bold'):
        report('REFERENCES_LABEL', 'References: is not fully bold.')
    for table in summary.tables:
        for index, row in enumerate(table.rows[1:], 1):
            if len(row.cells) < 2 or not row.cells[1].text.strip():
                report('SUMMARY_EMPTY', f'Nested summary row {index}: company cell is empty.')

    top_lines = reference_lines(summary)
    top = dict(top_lines)
    if len(top) != len(top_lines):
        report('REFERENCE_DUPLICATE', 'Summary reference numbers are duplicated.')
    for number, urls in top_lines:
        if len(urls) != 1:
            report('REFERENCE_URL', f'Summary Ref-{number}: expected one source URL.')
    first_use = []
    for index, row in enumerate(document.tables[1].rows, 1):
        cell = row.cells[1]
        paragraphs = comment_paragraphs(cell)
        if not paragraphs or not paragraph_text(paragraphs[0]).lstrip().startswith('[Comment:'):
            report('COMMENT_OPENING', f'Row {index}: missing opening [Comment: label.')
        else:
            opening = paragraphs[0]
            start = paragraph_text(opening).index('[Comment:')
            if not all(span_has(opening, start, start + len('[Comment:'), prop) for prop in ('bold', 'italic')):
                report('COMMENT_OPENING', f'Row {index}: [Comment: is not bold italic.')
        if not allow_inline and any(INLINE_REF.search(paragraph_text(p)) for p in paragraphs):
            report('INLINE_REFERENCE', f'Row {index}: move reference tags from the Comment to its source area.')
        closing = paragraphs[-1] if paragraphs else None
        text = paragraph_text(closing).rstrip() if closing is not None else ''
        match = HEREIN.search(text)
        if not match or not text.endswith(']'):
            report('HEREIN', f'Row {index}: missing complete Herein correspondence paragraph ending in ].')
        else:
            if not all(span_has(closing, len(text) - 1, len(text), prop) for prop in ('bold', 'italic')):
                report('COMMENT_CLOSING', f'Row {index}: final ] is not bold italic.')
            if require_emphasis:
                if not span_has(closing, *match.span('claim'), 'bold'):
                    report('CLAIM_EMPHASIS', f'Row {index}: the full quoted claim phrase must be bold.')
                start = match.start('mapping')
                if not span_has(closing, start, start + 1, 'bold'):
                    report('MAPPING_EMPHASIS', f'Row {index}: the corresponding element must begin with a bold phrase.')
                connector = text.index('corresponds to', match.start())
                if span_has(closing, connector, connector + len('corresponds to'), 'bold'):
                    report('CORRESPONDENCE_FRAME', f'Row {index}: keep the connector in the body style, rather than bolding all of Herein.')
            if template is not None and index <= len(template.tables[1].rows):
                reference_row = template.tables[1].rows[index - 1]
                expected = next(filter(None, (
                    HEREIN.search(paragraph_text(p))
                    for p in comment_paragraphs(reference_row.cells[1])
                )), None)
                same_left = ' '.join(row.cells[0].text.split()) == ' '.join(reference_row.cells[0].text.split())
                if expected and same_left and ' '.join(match.group('claim').split()) != ' '.join(expected.group('claim').split()):
                    report('CLAIM_PHRASE', f'Row {index}: the Herein claim phrase differs from the reference for the same left-column limitation.')
        if not image_first(evidence_events(cell)):
            report('EVIDENCE_ORDER', f'Row {index}: pair each screenshot with its own following URL; supplementary references may follow.')
        for number, urls in reference_lines(cell):
            if number not in first_use:
                first_use.append(number)
            if len(urls) != 1:
                report('REFERENCE_URL', f'Row {index}, Ref-{number}: expected one source URL.')
            if number not in top:
                report('REFERENCE_MISSING', f'Row {index}, Ref-{number}: missing from the summary.')
            elif urls != top[number]:
                report('REFERENCE_MISMATCH', f'Row {index}, Ref-{number}: URL differs from the summary.')
    if [number for number, _ in top_lines] != first_use:
        report('REFERENCE_ORDER', 'Summary references must cover the row-level sources in first-use order.')
    return issues


def save_repaired(document, source, output, changes):
    if source.resolve() == output.resolve() or (output.exists() and os.path.samefile(source, output)):
        raise UnsafeRepair('Choose a separate --output path; the input document is preserved.')
    content = source.read_bytes()
    if changes:
        rendered = BytesIO()
        document.save(rendered)
        part_name = str(document.part.partname).lstrip('/')
        with ZipFile(BytesIO(rendered.getvalue())) as modified:
            body = modified.read(part_name)
        rebuilt = BytesIO()
        with ZipFile(BytesIO(content)) as original, ZipFile(rebuilt, 'w') as result:
            result.comment = original.comment
            for info in original.infolist():
                result.writestr(info, body if info.filename == part_name else original.read(info.filename))
        content = rebuilt.getvalue()
    with tempfile.NamedTemporaryFile(dir=output.parent, prefix='.claim-chart-', suffix='.docx', delete=False) as temporary:
        temporary.write(content)
        temporary_name = Path(temporary.name)
    try:
        os.replace(temporary_name, output)
    finally:
        temporary_name.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('docx', type=Path)
    parser.add_argument('--check', action='store_true', help='Read-only validation (the default).')
    parser.add_argument('--template', type=Path, help='Reference DOCX for summary prose, inline citations and Herein emphasis/wording.')
    parser.add_argument('--fix-labels', action='store_true', help='Repair only References:, [Comment: and the final ].')
    parser.add_argument('--reorder-url-first', action='store_true', help='Explicitly convert a confirmed, uniform URL/image sequence to image/URL.')
    parser.add_argument('--output', type=Path, help='Separate output DOCX required for repairs; never overwrites the input.')
    args = parser.parse_args()
    repair = args.fix_labels or args.reorder_url_first
    if args.check and repair:
        parser.error('--check cannot be combined with repair options.')
    if bool(args.output) != bool(repair):
        parser.error('Repairs require --output; check-only runs do not write an output.')
    document = Document(args.docx)
    template = Document(args.template) if args.template else None
    changes = 0
    if repair:
        if not supported_layout(document):
            parser.error('Unsupported document layout; inspect and edit it manually.')
        try:
            if args.fix_labels:
                changes += format_references_label(document)
            for row in document.tables[1].rows:
                if args.fix_labels:
                    changes += format_comment_opening(row.cells[1])
                    changes += format_comment_closing(row.cells[1])
                if args.reorder_url_first:
                    changes += reorder_evidence(row.cells[1])
            save_repaired(document, args.docx, args.output, changes)
        except UnsafeRepair as error:
            parser.error(str(error))
        print(f'Saved {args.output}; repaired {changes} marker(s)/pair(s). Input preserved.')
    issues = collect_issues(document, template)
    if issues:
        print('\n'.join(map(str, issues)))
        return 1
    print('No supported claim-chart formatting issues found. Factual and rendered-page review are still required.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
