"""Synthetic DOCX regressions. No client documents or web evidence are fixtures."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile
import zlib

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import polish_claim_chart_docx as polish


SCRIPT = Path(polish.__file__)


def png(color):
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))
    header = struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header) + chunk(b'IDAT', zlib.compress(b'\0' + bytes(color))) + chunk(b'IEND', b'')


def hyperlink(paragraph, url, label=None):
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = label or url
    run.append(text)
    element.append(run)
    paragraph._p.append(element)
    return element


def reference(cell, number, fragment=''):
    paragraph = cell.add_paragraph(f'[Ref-{number}] ')
    hyperlink(paragraph, f'https://example.com/source-{number}' + fragment)
    return paragraph


def chart(order='image-first', broken_labels=False):
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(12)
    body_style = doc.styles.add_style('Comparison Body', WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = doc.styles['Normal']
    body_style.font.italic = True
    summary = doc.add_table(rows=1, cols=2)
    summary.autofit = False
    summary.cell(0, 0).text = 'Synthetic claim chart'
    top = summary.cell(0, 1)
    top.paragraphs[0].text = 'Example Supplier'
    label = top.add_paragraph().add_run('References:')
    label.bold = not broken_labels
    reference(top, 1)
    reference(top, 2)
    nested = top.add_table(rows=2, cols=2)
    nested.autofit = False
    nested.cell(0, 0).text = 'Claim Feature'
    nested.cell(0, 1).text = 'Corresponding Element'
    nested.cell(1, 0).text = 'Sensing module'
    nested.cell(1, 1).text = 'Example perception subsystem'
    main = doc.add_table(rows=1, cols=2)
    main.autofit = False
    main.cell(0, 0).text = 'a sensing module that collects road information;'
    cell = main.cell(0, 1)
    opening = cell.paragraphs[0]
    opening.style = body_style
    opening.paragraph_format.keep_with_next = True
    opening.paragraph_format.space_after = Pt(7)
    if broken_labels:
        opening.add_run('[Com').bold = False
        opening.add_run('ment: The published system has ').bold = False
    else:
        marker = opening.add_run('[Comment: ')
        marker.bold = marker.italic = True
        opening.add_run('The published system has ')
    opening.add_run('selective emphasis').bold = True
    opening.add_run(' and a ')
    hyperlink(opening, 'https://example.com/explanation', 'technical description')
    opening.add_run('.')
    herein = cell.add_paragraph()
    herein.style = body_style
    herein.add_run('Herein, “')
    herein.add_run('a sensing module that collects road information').bold = True
    herein.add_run('” corresponds to ')
    herein.add_run('the Example perception subsystem').bold = True
    if broken_labels:
        herein.add_run(', which senses traffic.]   ')
    else:
        herein.add_run(', which senses traffic.')
        marker = herein.add_run(']')
        marker.bold = marker.italic = True
    for number in (1, 2):
        if order == 'url-first':
            reference(cell, number)
        image = cell.add_paragraph()
        shape = image.add_run().add_picture(BytesIO(png((number * 80, 0, 0))), width=Inches(1))
        shape._inline.docPr.set('descr', f'Synthetic image for source {number}')
        if order == 'image-first':
            reference(cell, number)
    doc.sections[0].footer.paragraphs[0].text = 'Synthetic footer, retained verbatim'
    return doc


def codes(document, template=None):
    return {issue.code for issue in polish.collect_issues(document, template)}


def protected_geometry_and_left_cells(document):
    return (
        [etree.tostring(element) for element in document.element.body.xpath('.//w:tblW | .//w:tblGrid | .//w:tcW')],
        [row.xpath('./w:tc')[0].xml for table in polish.body_tables(document) for row in table.xpath('./w:tr')],
    )


def without_table_layouts(document):
    copy = deepcopy(document)
    for layout in copy.element.body.xpath('.//w:tblPr/w:tblLayout'):
        layout.getparent().remove(layout)
    return copy.element.xml


class TableLayoutTests(unittest.TestCase):
    def test_checker_finds_missing_and_autofit_nested_layouts(self):
        doc = chart()
        tables = polish.body_tables(doc)
        self.assertEqual(len(tables), 3)
        missing = tables[0].tblPr.tblLayout
        missing.getparent().remove(missing)
        tables[1].tblPr.tblLayout.set(qn('w:type'), 'autofit')
        before = doc.element.xml
        issues = [issue for issue in polish.collect_issues(doc) if issue.code == 'TABLE_LAYOUT']
        self.assertEqual(len(issues), 2)
        self.assertIn('Table 1', issues[0].message)
        self.assertIn('missing', issues[0].message)
        self.assertIn('Table 2', issues[1].message)
        self.assertIn('autofit', issues[1].message)
        self.assertEqual(doc.element.xml, before)

    def test_long_urls_do_not_justify_changing_widths_or_content(self):
        doc = chart()
        long_url = 'https://example.com/' + 'continuous_path_segment_' * 100
        hyperlink(doc.tables[0].cell(0, 1).add_paragraph(), long_url)
        hyperlink(doc.tables[1].cell(0, 1).add_paragraph(), long_url)
        tables = polish.body_tables(doc)
        for index, table in enumerate(tables):
            table.tblPr.find(qn('w:tblW')).set(qn('w:w'), str(12000 - index * 1000))
            table.tblPr.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
            for column, grid in enumerate(table.tblGrid):
                grid.set(qn('w:w'), str(3000 + column * 4000 - index * 500))
            for cell_index, width in enumerate(table.xpath('./w:tr/w:tc/w:tcPr/w:tcW')):
                width.set(qn('w:w'), str(3100 + cell_index * 100 + index * 10))
            layout = table.tblPr.tblLayout
            if index == 0:
                layout.getparent().remove(layout)
            else:
                layout.set(qn('w:type'), 'autofit')
        protected = protected_geometry_and_left_cells(doc)
        other_content = without_table_layouts(doc)
        self.assertEqual(polish.fix_table_layout(doc), 3)
        self.assertEqual(protected_geometry_and_left_cells(doc), protected)
        self.assertEqual(without_table_layouts(doc), other_content)
        self.assertNotIn('TABLE_LAYOUT', codes(doc))
        before_second = doc.element.xml
        self.assertEqual(polish.fix_table_layout(doc), 0)
        self.assertEqual(doc.element.xml, before_second)

    def test_absent_table_properties_are_created_without_changing_grid_or_cells(self):
        doc = chart()
        table = polish.body_tables(doc)[1]
        table.remove(table.tblPr)
        protected = protected_geometry_and_left_cells(doc)
        self.assertIn('TABLE_LAYOUT', codes(doc))
        self.assertEqual(polish.fix_table_layout(doc), 1)
        self.assertEqual(protected_geometry_and_left_cells(doc), protected)
        self.assertEqual(table[0].tag, qn('w:tblPr'))
        self.assertEqual(table.tblPr.tblLayout.get(qn('w:type')), 'fixed')

    def test_duplicate_properties_are_rejected_before_any_table_is_mutated(self):
        doc = chart()
        tables = polish.body_tables(doc)
        tables[0].tblPr.tblLayout.set(qn('w:type'), 'autofit')
        tables[-1].tblPr.append(deepcopy(tables[-1].tblPr.tblLayout))
        original = doc.element.xml
        self.assertIn('TABLE_LAYOUT', codes(doc))
        with self.assertRaises(polish.UnsafeRepair):
            polish.fix_table_layout(doc)
        self.assertEqual(doc.element.xml, original)

    def test_cli_fix_is_explicit_separate_and_preserves_unrelated_package_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output, repeated = [Path(directory) / name for name in ('source.docx', 'output.docx', 'repeated.docx')]
            doc = chart()
            for table in polish.body_tables(doc):
                table.tblPr.tblLayout.set(qn('w:type'), 'autofit')
            doc.save(source)
            with ZipFile(source, 'a') as archive:
                archive.comment = b'Synthetic archive metadata.'
                archive.writestr('opaque-fixture.bin', b'Preserve this unrelated package part.')
            original = source.read_bytes()
            result = subprocess.run([sys.executable, str(SCRIPT), str(source)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
            self.assertEqual(result.stdout.count('[TABLE_LAYOUT]'), 3)
            self.assertEqual(source.read_bytes(), original)
            self.assertFalse(output.exists())
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), '--fix-table-layout', '--output', str(output)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(source.read_bytes(), original)
            with ZipFile(source) as old, ZipFile(output) as new:
                self.assertEqual(old.namelist(), new.namelist())
                self.assertEqual(old.comment, new.comment)
                self.assertEqual([name for name in old.namelist() if old.read(name) != new.read(name)], ['word/document.xml'])
            repaired = Document(output)
            self.assertEqual(protected_geometry_and_left_cells(doc), protected_geometry_and_left_cells(repaired))
            self.assertEqual(without_table_layouts(doc), without_table_layouts(repaired))
            self.assertEqual(codes(repaired), set())
            result = subprocess.run([sys.executable, str(SCRIPT), str(output), '--fix-table-layout', '--output', str(repeated)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(repeated.read_bytes(), output.read_bytes())

    def test_cli_layout_repair_rejects_check_mode_missing_output_and_in_place_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.docx'
            chart().save(source)
            original = source.read_bytes()
            for extra in (['--fix-table-layout'], ['--fix-table-layout', '--output', str(source)], ['--check', '--fix-table-layout', '--output', str(Path(directory) / 'output.docx')]):
                result = subprocess.run([sys.executable, str(SCRIPT), str(source), *extra], capture_output=True, text=True)
                self.assertEqual(result.returncode, 2)
                self.assertEqual(source.read_bytes(), original)


class EvidenceTests(unittest.TestCase):
    def test_already_correct_pairs_are_unchanged_on_repeated_repairs(self):
        doc = chart()
        cell = doc.tables[1].cell(0, 1)
        original = cell._tc.xml
        self.assertEqual(polish.reorder_evidence(cell), 0)
        self.assertEqual(polish.reorder_evidence(cell), 0)
        self.assertEqual(cell._tc.xml, original)
        self.assertEqual(codes(doc), set())

    def test_explicit_url_first_repair_keeps_source_image_identity(self):
        doc = chart('url-first')
        cell = doc.tables[1].cell(0, 1)
        events = polish.evidence_events(cell)
        before = [(events[i][1]._p.xml, events[i + 1][1]._p.xml) for i in (0, 2)]
        self.assertIn('EVIDENCE_ORDER', codes(doc))
        self.assertEqual(polish.reorder_evidence(cell), 2)
        events = polish.evidence_events(cell)
        after = [(events[i + 1][1]._p.xml, events[i][1]._p.xml) for i in (0, 2)]
        self.assertEqual(before, after)
        self.assertEqual(polish.reorder_evidence(cell), 0)
        self.assertEqual(codes(doc), set())

    def test_mixed_evidence_is_rejected_without_moving_anything(self):
        doc = chart('url-first')
        cell = doc.tables[1].cell(0, 1)
        reference(cell, 3)
        original = cell._tc.xml
        with self.assertRaises(polish.UnsafeRepair):
            polish.reorder_evidence(cell)
        self.assertEqual(original, cell._tc.xml)

    def test_caption_between_url_and_image_requires_manual_pairing(self):
        doc = chart('url-first')
        cell = doc.tables[1].cell(0, 1)
        events = polish.evidence_events(cell)
        caption = cell.add_paragraph('Caption belongs to an evidence block.')
        events[1][1]._p.addprevious(caption._p)
        original = cell._tc.xml
        with self.assertRaises(polish.UnsafeRepair):
            polish.reorder_evidence(cell)
        self.assertEqual(original, cell._tc.xml)

    def test_supplementary_references_and_pdf_page_fragments_are_valid(self):
        doc = chart()
        reference(doc.tables[0].cell(0, 1), 3)
        reference(doc.tables[1].cell(0, 1), 3, '#page=18')
        self.assertEqual(codes(doc), set())

    def test_added_caption_and_dates_are_rejected_when_template_has_none(self):
        reference_doc = chart()
        doc = chart()
        cell = doc.tables[1].cell(0, 1)
        cell.add_paragraph('Official announcement; excerpt identifying the subsystem. Published 2026-01-01.')
        self.assertIn('EVIDENCE_PROSE', codes(doc, reference_doc))
        self.assertIn('EVIDENCE_PROSE', codes(doc))

    def test_prose_appended_to_url_line_is_also_reported(self):
        doc = chart()
        cell = doc.tables[1].cell(0, 1)
        polish.evidence_events(cell)[1][1].add_run(' Accessed September 5, 2026.')
        self.assertIn('EVIDENCE_PROSE', codes(doc, chart()))

    def test_template_caption_convention_is_respected(self):
        template = chart()
        template.tables[1].cell(0, 1).add_paragraph('Source description required by this reference.')
        doc = chart()
        doc.tables[1].cell(0, 1).add_paragraph('Source description for the assessed product.')
        self.assertEqual(codes(doc, template), set())


class FormattingTests(unittest.TestCase):
    def test_marker_repairs_preserve_body_typography_hyperlinks_and_properties(self):
        doc = chart(broken_labels=True)
        cell = doc.tables[1].cell(0, 1)
        paragraphs = polish.comment_paragraphs(cell)
        text_before = [polish.paragraph_text(p) for p in paragraphs]
        opening = paragraphs[0]
        ppr = opening._p.pPr.xml
        link = opening._p.xpath('.//w:hyperlink')[0].xml
        emphasis = next(r for r in opening.runs if r.text == 'selective emphasis')
        emphasis_xml = emphasis._r.xml
        left = doc.tables[1].cell(0, 0)._tc.xml
        self.assertEqual(polish.format_references_label(doc), 1)
        self.assertEqual(polish.format_comment_opening(cell), 1)
        self.assertEqual(polish.format_comment_closing(cell), 1)
        self.assertEqual([polish.paragraph_text(p) for p in polish.comment_paragraphs(cell)], text_before)
        self.assertEqual(opening._p.pPr.xml, ppr)
        self.assertEqual(opening._p.xpath('.//w:hyperlink')[0].xml, link)
        self.assertEqual(emphasis._r.xml, emphasis_xml)
        self.assertEqual(doc.tables[1].cell(0, 0)._tc.xml, left)
        body = next(r for r in opening.runs if 'The published system' in r.text)
        self.assertFalse(polish.effective(body, opening, 'bold'))
        self.assertTrue(polish.effective(body, opening, 'italic'))
        self.assertIsNone(body.font.name)
        self.assertEqual(doc.styles['Normal'].font.name, 'Calibri')
        self.assertEqual(doc.styles['Normal'].font.size.pt, 12)
        self.assertEqual(codes(doc), set())
        before_second = doc.element.xml
        self.assertEqual(polish.format_references_label(doc), 0)
        self.assertEqual(polish.format_comment_opening(cell), 0)
        self.assertEqual(polish.format_comment_closing(cell), 0)
        self.assertEqual(doc.element.xml, before_second)

    def test_character_style_inheritance_is_recognized(self):
        doc = chart()
        style = doc.styles.add_style('Comment Marker', WD_STYLE_TYPE.CHARACTER)
        style.font.bold = style.font.italic = True
        opening = doc.tables[1].cell(0, 1).paragraphs[0]
        opening.runs[0].style = style
        opening.runs[0].bold = opening.runs[0].italic = None
        before = opening._p.xml
        self.assertEqual(polish.format_comment_opening(doc.tables[1].cell(0, 1)), 0)
        self.assertEqual(opening._p.xml, before)
        self.assertNotIn('COMMENT_OPENING', codes(doc))

    def test_three_reported_template_regressions_are_detected(self):
        reference_doc = chart()
        doc = chart()
        doc.tables[0].cell(0, 1).add_paragraph('Comparison as of a new date; extra scope prose.')
        cell = doc.tables[1].cell(0, 1)
        cell.paragraphs[0].add_run(' [Ref-1; Ref-2]')
        herein = polish.comment_paragraphs(cell)[-1]
        for run in herein.runs:
            if run.text not in (']',):
                run.bold = False
        self.assertTrue({'SUMMARY_PROSE', 'INLINE_REFERENCE', 'CLAIM_EMPHASIS', 'MAPPING_EMPHASIS'} <= codes(doc, reference_doc))

    def test_template_can_explicitly_use_prose_inline_refs_and_plain_correspondence(self):
        doc = chart()
        doc.tables[0].cell(0, 1).add_paragraph('Scope required by this reference.')
        cell = doc.tables[1].cell(0, 1)
        cell.paragraphs[0].add_run(' [Ref-1]')
        for run in polish.comment_paragraphs(cell)[-1].runs:
            if run.text != ']':
                run.bold = False
        self.assertEqual(codes(doc, deepcopy(doc)), set())

    def test_entire_herein_bold_is_not_a_substitute_for_two_spans(self):
        doc = chart()
        for run in polish.comment_paragraphs(doc.tables[1].cell(0, 1))[-1].runs:
            run.bold = True
        self.assertIn('CORRESPONDENCE_FRAME', codes(doc))

    def test_abbreviated_claim_is_detected_without_a_reference(self):
        doc = chart()
        herein = polish.comment_paragraphs(doc.tables[1].cell(0, 1))[-1]
        herein.runs[1].text = 'a sensing module ... road information'
        self.assertIn('CLAIM_PHRASE', codes(doc))

    def test_complete_unnumbered_quote_is_checked_against_left_not_short_template_quote(self):
        for enumeration in ('1. ', '1) ', '2) ', '(3) '):
            with self.subTest(enumeration=enumeration):
                doc = chart()
                left = doc.tables[1].cell(0, 0)
                left.text = enumeration + left.text
                reference_doc = deepcopy(doc)
                reference_quote = polish.comment_paragraphs(reference_doc.tables[1].cell(0, 1))[-1]
                reference_quote.runs[1].text = enumeration + 'a sensing module'
                protected = protected_geometry_and_left_cells(doc)
                self.assertEqual(codes(doc, reference_doc), set())
                self.assertEqual(protected_geometry_and_left_cells(doc), protected)

    def test_numbered_right_quote_is_reported_even_when_the_substantive_text_is_complete(self):
        for enumeration in ('1. ', '1) ', '2) ', '(3) '):
            with self.subTest(enumeration=enumeration):
                doc = chart()
                left = doc.tables[1].cell(0, 0)
                left.text = enumeration + left.text
                herein = polish.comment_paragraphs(doc.tables[1].cell(0, 1))[-1]
                herein.runs[1].text = enumeration + herein.runs[1].text
                self.assertIn('CLAIM_NUMBERING', codes(doc))
                self.assertNotIn('CLAIM_PHRASE', codes(doc))

    def test_missing_functional_condition_is_reported_even_if_template_also_omits_it(self):
        doc = chart()
        doc.tables[1].cell(0, 0).text = '2) a sensing module that collects road information and uses a backup plan if a sensor fails; and'
        self.assertIn('CLAIM_PHRASE', codes(doc, deepcopy(doc)))
        herein = polish.comment_paragraphs(doc.tables[1].cell(0, 1))[-1]
        herein.runs[1].text = 'a sensing module that collects road information and uses a backup plan if a sensor fails; and'
        self.assertEqual(codes(doc), set())

    def test_substantive_numeric_values_are_not_treated_as_enumeration(self):
        doc = chart()
        phrase = '1.5 GHz communication circuitry for a 2) protocol field'
        doc.tables[1].cell(0, 0).text = phrase + ';'
        herein = polish.comment_paragraphs(doc.tables[1].cell(0, 1))[-1]
        herein.runs[1].text = phrase
        self.assertEqual(codes(doc), set())

    def test_summary_status_does_not_substitute_for_an_actual_technology_name(self):
        doc = chart()
        summary = doc.tables[0].cell(0, 1).tables[0].cell(1, 1)
        for text in ('Unconfirmed', 'Proposed Example recovery control', 'Example subsystem (candidate)'):
            with self.subTest(text=text):
                summary.text = text
                self.assertIn('SUMMARY_STATUS', codes(doc))
        summary.text = 'Example stop-and-replan recovery control'
        self.assertEqual(codes(doc), set())

    def test_gap_analysis_cannot_replace_the_required_herein_in_partial_mapping(self):
        doc = chart()
        cell = doc.tables[1].cell(0, 1)
        herein = polish.comment_paragraphs(cell)[-1]
        herein.text = 'No complete correspondence is established for this requirement.]'
        self.assertIn('HEREIN', codes(doc))

    def test_empty_nested_summary_and_missing_source_are_reported(self):
        doc = chart()
        doc.tables[0].cell(0, 1).tables[0].cell(1, 1).text = ''
        reference(doc.tables[1].cell(0, 1), 3)
        self.assertTrue({'SUMMARY_EMPTY', 'REFERENCE_MISSING', 'REFERENCE_ORDER'} <= codes(doc))

    def test_source_number_cannot_silently_point_to_a_different_url(self):
        doc = chart()
        cell = doc.tables[1].cell(0, 1)
        paragraph = polish.evidence_events(cell)[1][1]
        link = paragraph._p.xpath('.//w:hyperlink')[0]
        # Keep the top line intact and change only the row's hyperlink target.
        link.set(qn('r:id'), paragraph.part.relate_to('https://example.com/third', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
        self.assertIn('REFERENCE_MISMATCH', codes(doc))

    def test_missing_hereins_do_not_misclassify_source_lines_as_inline_tags(self):
        doc = chart()
        cell = doc.tables[1].cell(0, 1)
        herein = polish.comment_paragraphs(cell)[-1]
        herein._p.getparent().remove(herein._p)
        self.assertIn('HEREIN', codes(doc))
        self.assertNotIn('INLINE_REFERENCE', codes(doc))

    def test_unsupported_layout_returns_diagnostic(self):
        doc = Document()
        doc.add_table(rows=1, cols=1)
        doc.add_table(rows=1, cols=1)
        self.assertEqual(codes(doc), {'LAYOUT'})


class FileTests(unittest.TestCase):
    def test_default_cli_is_read_only(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.docx'
            chart().save(source)
            original = source.read_bytes()
            result = subprocess.run([sys.executable, str(SCRIPT), str(source)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(source.read_bytes(), original)

    def test_repaired_output_preserves_original_and_all_other_package_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = [Path(directory) / name for name in ('source.docx', 'output.docx')]
            chart(broken_labels=True).save(source)
            with ZipFile(source, 'a') as archive:
                archive.writestr('opaque-fixture.bin', b'Preserve this unrelated package part.')
            original = source.read_bytes()
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), '--fix-labels', '--output', str(output)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(source.read_bytes(), original)
            with ZipFile(source) as old, ZipFile(output) as new:
                self.assertEqual(old.namelist(), new.namelist())
                self.assertEqual([name for name in old.namelist() if old.read(name) != new.read(name)], ['word/document.xml'])
            self.assertEqual(codes(Document(output)), set())

    def test_repairs_require_separate_output_and_reject_in_place_path(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.docx'
            chart().save(source)
            original = source.read_bytes()
            for extra in (['--fix-labels'], ['--fix-labels', '--output', str(source)]):
                result = subprocess.run([sys.executable, str(SCRIPT), str(source), *extra], capture_output=True, text=True)
                self.assertEqual(result.returncode, 2)
                self.assertEqual(source.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
